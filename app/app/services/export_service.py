import csv
import io
import zipfile
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Tuple
from app.db.connection import db

class ExportService:
    @staticmethod
    def _create_styled_workbook(sheet_title: str, headers: List[str], rows: List[List[Any]]) -> io.BytesIO:
        """Creates a cleanly formatted Excel spreadsheet using openpyxl."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]  # Excel max sheet title is 31 chars

        # Header styling
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="8B4513", end_color="8B4513", fill_type="solid")  # Brand warm amber/brown
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E0E0E0'),
            right=Side(style='thin', color='E0E0E0'),
            top=Side(style='thin', color='E0E0E0'),
            bottom=Side(style='thin', color='E0E0E0')
        )

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Data rows
        data_font = Font(name="Calibri", size=10)
        data_align = Alignment(vertical="center")
        
        for row_data in rows:
            formatted_row = []
            for val in row_data:
                if isinstance(val, (datetime.datetime, datetime.date)):
                    formatted_row.append(val.strftime("%Y-%m-%d %H:%M:%S"))
                elif val is None:
                    formatted_row.append("")
                else:
                    formatted_row.append(str(val))
            ws.append(formatted_row)

        # Apply borders and auto-fit column widths
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    @staticmethod
    def _create_csv(headers: List[str], rows: List[List[Any]]) -> io.StringIO:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for r in rows:
            writer.writerow([str(v) if v is not None else "" for v in r])
        output.seek(0)
        return output

    # ----------------- Products -----------------
    @staticmethod
    def get_products_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["ID", "Name", "SKU", "Category", "MOQ", "Material", "Sizes", "Colours", "Availability", "Featured", "New Arrival", "Archived", "Created At"]
        items = db.fetch_all("""
            SELECT p.id, p.name, p.sku, c.name as category_name, p.moq, p.material, p.size_info, p.available_colours,
                   p.availability, p.is_featured, p.is_new_arrival, p.is_archived, p.created_at
            FROM products p
            LEFT JOIN categories c ON p.category_id = c.id
            ORDER BY p.id ASC
        """)
        rows = [[
            i["id"], i["name"], i["sku"], i["category_name"], i["moq"], i["material"], i["size_info"],
            i["available_colours"], i["availability"], i["is_featured"], i["is_new_arrival"], i["is_archived"], i["created_at"]
        ] for i in items]
        return headers, rows

    # ----------------- Customers -----------------
    @staticmethod
    def get_customers_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["Customer ID", "Google ID", "Name", "Email", "Mobile", "Status", "Created At", "Last Login"]
        items = db.fetch_all("SELECT id, google_id, name, email, mobile, status, created_at, last_login_at FROM customers ORDER BY created_at DESC")
        rows = [[
            i["id"], i["google_id"], i["name"], i["email"], i["mobile"], i["status"], i["created_at"], i["last_login_at"]
        ] for i in items]
        return headers, rows

    # ----------------- Enquiries -----------------
    @staticmethod
    def get_enquiries_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["Enquiry Ref", "Date", "Customer Name", "Mobile", "Email", "Status", "Total Items", "Notes"]
        items = db.fetch_all("""
            SELECT e.enquiry_reference, e.created_at, e.customer_name, e.customer_mobile, e.customer_email,
                   e.status, e.notes,
                   (SELECT count(*) FROM enquiry_items WHERE enquiry_id = e.id) as item_count
            FROM enquiries e
            ORDER BY e.created_at DESC
        """)
        rows = [[
            i["enquiry_reference"], i["created_at"], i["customer_name"], i["customer_mobile"], i["customer_email"],
            i["status"], i["item_count"], i["notes"]
        ] for i in items]
        return headers, rows

    # ----------------- Enquiry Items Snapshot -----------------
    @staticmethod
    def get_enquiry_items_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["Enquiry Ref", "Product Name", "SKU", "Category", "Size", "Colour", "Quantity", "MOQ At Enquiry"]
        items = db.fetch_all("""
            SELECT e.enquiry_reference, ei.product_name, ei.sku, ei.category_name,
                   ei.selected_size, ei.selected_colour, ei.quantity, ei.moq_at_enquiry
            FROM enquiry_items ei
            JOIN enquiries e ON ei.enquiry_id = e.id
            ORDER BY e.created_at DESC, ei.id ASC
        """)
        rows = [[
            i["enquiry_reference"], i["product_name"], i["sku"], i["category_name"],
            i["selected_size"], i["selected_colour"], i["quantity"], i["moq_at_enquiry"]
        ] for i in items]
        return headers, rows

    # ----------------- Categories -----------------
    @staticmethod
    def get_categories_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["ID", "Name", "Slug", "Description", "Display Order", "Active"]
        items = db.fetch_all("SELECT id, name, slug, description, display_order, is_active FROM categories ORDER BY display_order ASC")
        rows = [[i["id"], i["name"], i["slug"], i["description"], i["display_order"], i["is_active"]] for i in items]
        return headers, rows

    # ----------------- Website Settings -----------------
    @staticmethod
    def get_settings_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["Setting Key", "Value", "Type", "Description", "Updated At"]
        items = db.fetch_all("SELECT key, value, type, description, updated_at FROM website_settings ORDER BY key ASC")
        rows = [[i["key"], i["value"], i["type"], i["description"], i["updated_at"]] for i in items]
        return headers, rows

    # ----------------- Social Links -----------------
    @staticmethod
    def get_social_links_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["ID", "Platform", "Name", "URL", "Icon", "Display Order", "Active"]
        items = db.fetch_all("SELECT id, platform, name, url, icon, display_order, is_enabled FROM social_links ORDER BY display_order ASC")
        rows = [[i["id"], i["platform"], i["name"], i["url"], i["icon"], i["display_order"], i["is_enabled"]] for i in items]
        return headers, rows

    # ----------------- FAQs -----------------
    @staticmethod
    def get_faqs_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["ID", "Question", "Answer", "Display Order", "Active"]
        items = db.fetch_all("SELECT id, question, answer, display_order, is_active FROM faqs ORDER BY display_order ASC")
        rows = [[i["id"], i["question"], i["answer"], i["display_order"], i["is_active"]] for i in items]
        return headers, rows

    # ----------------- Media Manifest -----------------
    @staticmethod
    def get_media_manifest_data() -> Tuple[List[str], List[List[Any]]]:
        headers = ["ID", "File URL", "File Name", "MIME Type", "Size Bytes", "Entity Type", "Entity ID", "Created At"]
        items = db.fetch_all("SELECT * FROM media_manifest ORDER BY created_at DESC")
        rows = [[i["id"], i["file_url"], i["file_name"], i["mime_type"], i["file_size"], i["entity_type"], i["entity_id"], i["created_at"]] for i in items]
        return headers, rows

    # ----------------- Main Export Dispatcher -----------------
    @classmethod
    def export_table(cls, table_name: str, format_type: str = "xlsx") -> Tuple[bytes, str, str]:
        """Returns (content_bytes, filename, media_type)."""
        data_map = {
            "products": (cls.get_products_data, "products"),
            "customers": (cls.get_customers_data, "customers"),
            "enquiries": (cls.get_enquiries_data, "enquiries"),
            "enquiry_items": (cls.get_enquiry_items_data, "enquiry_items"),
            "categories": (cls.get_categories_data, "categories"),
            "settings": (cls.get_settings_data, "website_settings"),
            "social_links": (cls.get_social_links_data, "social_links"),
            "faq": (cls.get_faqs_data, "faq"),
            "media_manifest": (cls.get_media_manifest_data, "media_manifest")
        }
        if table_name not in data_map:
            raise ValueError(f"Unknown export target: {table_name}")

        getter, title = data_map[table_name]
        headers, rows = getter()
        date_str = datetime.datetime.now().strftime("%Y-%m-%d")

        if format_type.lower() == "csv":
            csv_io = cls._create_csv(headers, rows)
            return csv_io.getvalue().encode("utf-8"), f"{title}_{date_str}.csv", "text/csv"
        else:
            buf = cls._create_styled_workbook(title.title(), headers, rows)
            return buf.getvalue(), f"{title}_{date_str}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @classmethod
    def export_everything_zip(cls) -> Tuple[bytes, str]:
        """Generates Shree-Shyam-Toys-Backup-YYYY-MM-DD.zip with all XLSX files."""
        date_str = datetime.datetime.now().strftime("%Y-%m-%d")
        zip_filename = f"Shree-Shyam-Toys-Backup-{date_str}.zip"

        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
            tables = [
                ("products.xlsx", cls.get_products_data, "Products"),
                ("categories.xlsx", cls.get_categories_data, "Categories"),
                ("customers.xlsx", cls.get_customers_data, "Customers"),
                ("enquiries.xlsx", cls.get_enquiries_data, "Enquiries"),
                ("enquiry_items.xlsx", cls.get_enquiry_items_data, "Enquiry Items"),
                ("website_settings.xlsx", cls.get_settings_data, "Website Settings"),
                ("social_links.xlsx", cls.get_social_links_data, "Social Links"),
                ("faq.xlsx", cls.get_faqs_data, "FAQ"),
                ("media_manifest.xlsx", cls.get_media_manifest_data, "Media Manifest")
            ]
            for file_in_zip, getter, sheet_title in tables:
                headers, rows = getter()
                wb_buf = cls._create_styled_workbook(sheet_title, headers, rows)
                zf.writestr(file_in_zip, wb_buf.getvalue())

        zip_buf.seek(0)
        return zip_buf.getvalue(), zip_filename
